import argparse
import os
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

# =========================
# 指标定义
# =========================
METRICS = [
    "Request goodput (req/s)",
    "Output token goodput (tok/s)",
    "Total token goodput (tok/s)",
    "Mean TTFT (ms)",
    "P99 TTFT (ms)",
    "Mean TPOT (ms)",
    "P99 TPOT (ms)",
    "Mean E2EL (ms)",
    "P99 E2EL (ms)",
]

UP_METRICS = {
    "Request goodput (req/s)",
    "Output token goodput (tok/s)",
    "Total token goodput (tok/s)",
}

DOWN_METRICS = {
    "Mean TTFT (ms)",
    "P99 TTFT (ms)",
    "Mean TPOT (ms)",
    "P99 TPOT (ms)",
    "Mean E2EL (ms)",
    "P99 E2EL (ms)",
}

# =========================
# 从路径解析 method
# =========================
def extract_method_name(filepath: str) -> str:
    parts = filepath.split(os.sep)
    for p in parts:
        match = re.search(r"_test_([a-zA-Z0-9]+)", p)
        if match:
            return match.group(1)
    raise ValueError(f"Cannot extract method name from path: {filepath}")

# =========================
# 日志解析
# =========================
def parse_benchmark_log(filepath: str) -> dict:
    with open(filepath, "r", encoding="utf-8") as f:
        text = f.read()

    data = {}
    for metric in METRICS:
        pattern = re.escape(metric) + r":\s+([0-9.]+)"
        match = re.search(pattern, text)
        if not match:
            raise ValueError(f"Metric '{metric}' not found in {filepath}")
        data[metric] = float(match.group(1))
    return data

# =========================
# 主流程
# =========================
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--logs", nargs="+", required=True)
    parser.add_argument("--baseline", type=str, default=None)
    parser.add_argument("--outdir", type=str, default="./benchmark_results")
    args = parser.parse_args()

    os.makedirs(args.outdir, exist_ok=True)

    # -------------------------
    # 解析日志
    # -------------------------
    records = {}
    for path in args.logs:
        method = extract_method_name(path)
        if method in records:
            raise ValueError(f"Duplicate method detected: {method}")
        records[method] = parse_benchmark_log(path)

    df_abs = pd.DataFrame.from_dict(records, orient="index")[METRICS]

    # -------------------------
    # baseline 选择
    # -------------------------
    baseline_method = args.baseline or df_abs.index[0]
    if baseline_method not in df_abs.index:
        raise ValueError(f"Baseline '{baseline_method}' not found in {list(df_abs.index)}")

    baseline = df_abs.loc[baseline_method]

    # -------------------------
    # 相对 baseline 百分比变化
    # -------------------------
    df_delta = df_abs.copy()
    for col in METRICS:
        df_delta[col] = (df_abs[col] - baseline[col]) / baseline[col] * 100
    df_delta.loc[baseline_method, :] = 0.0

    # -------------------------
    # 保存 CSV
    # -------------------------
    abs_csv = os.path.join(args.outdir, "absolute_metrics.csv")
    rel_csv = os.path.join(args.outdir, "relative_change_vs_baseline.csv")

    df_abs.to_csv(abs_csv)
    df_delta.to_csv(rel_csv)

    # -------------------------
    # 生成带百分比标注的 Excel
    # -------------------------
    df_out = df_abs.copy()
    for method in df_abs.index:
        for col in METRICS:
            if method == baseline_method:
                df_out.loc[method, col] = f"{df_abs.loc[method, col]:.2f}"
            else:
                delta = df_delta.loc[method, col]
                df_out.loc[method, col] = f"{df_abs.loc[method, col]:.2f} ({delta:+.2f}%)"

    xlsx_path = os.path.join(args.outdir, "benchmark_metrics_annotated.xlsx")
    df_out.to_excel(xlsx_path)

    # -------------------------
    # Excel 红绿字体标注
    # -------------------------
    wb = load_workbook(xlsx_path)
    ws = wb.active

    green_font = Font(color="008000")
    red_font = Font(color="FF0000")

    for row_idx, method in enumerate(df_abs.index, start=2):
        if method == baseline_method:
            continue
        for col_idx, col in enumerate(METRICS, start=2):
            delta = df_delta.loc[method, col]
            cell = ws.cell(row=row_idx, column=col_idx)

            if col in UP_METRICS:
                cell.font = green_font if delta > 0 else red_font
            elif col in DOWN_METRICS:
                cell.font = green_font if delta < 0 else red_font

    wb.save(xlsx_path)

    # -------------------------
    # 完成提示
    # -------------------------
    print("Results saved:")
    print(f"  - {abs_csv}")
    print(f"  - {rel_csv}")
    print(f"  - {xlsx_path}")

if __name__ == "__main__":
    main()
